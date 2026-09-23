"""离线自检：不需要联网，用来验证依赖安装是否正常、解析与导出是否可用。

用法：

    python selftest.py

它会用几个模拟的 TikTok Shop 商品卡片 HTML 走一遍完整流程
（解析 → 去重 → 排名 → 写出 Excel → 读回校验）。
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

import config
from config import category_page_url, list_categories, resolve_category
from excel_export import export_to_excel
from page_parser import (
    assign_ranks,
    build_product_record,
    dedupe_records,
    extract_category_from_pdp,
    extract_discount,
    extract_prices,
    extract_product_description,
    extract_rating,
    extract_sold,
    is_meaningful,
    to_count,
)

# ---------------------------------------------------------------------------
# 模拟卡片：故意做成三种不同写法，覆盖常见 DOM 变体
# ---------------------------------------------------------------------------
CARDS = [
    # 变体 1：标准写法，有划线原价
    """
    <div class="product-card">
      <a href="/us/view/product/1729384756">
        <img src="https://p16-oec.example.com/img/1.jpg" alt="Wireless Earbuds Pro with Charging Case">
      </a>
      <a class="shop-name" href="/shop/12345">SoundCore Official Store</a>
      <div class="price"><span>$23.99</span><del>$39.99</del></div>
      <div class="meta">4.7★ &nbsp; 12.5K sold</div>
      <div class="badge">40% off</div>
    </div>
    """,
    # 变体 2：懒加载图片 + 只有一个价格 + Sold 前缀
    """
    <div class="ranking-item">
      <img data-src="https://p16-oec.example.com/img/2.jpg" alt="Stainless Steel Insulated Water Bottle 32oz">
      <a href="https://shop.tiktok.com/us/product/99887766" title="Stainless Steel Insulated Water Bottle 32oz">
        Stainless Steel Insulated Water Bottle 32oz
      </a>
      <span class="seller">HydroPure</span>
      <span>$15.49</span>
      <span>Sold 3.4K</span>
    </div>
    """,
    # 变体 3：没有 <del>，用 class 表示原价；评分写成 4.5 (2,314)
    """
    <li class="goods-card">
      <a href="/product/44556677">
        <img srcset="https://p16-oec.example.com/img/3_small.jpg 1x, https://p16-oec.example.com/img/3.jpg 2x"
             alt="LED Ring Light 18 inch with Tripod Stand">
      </a>
      <p class="product-title">LED Ring Light 18 inch with Tripod Stand</p>
      <span class="brand-name">GlowUp Studio</span>
      <span class="current-price">$42.00</span>
      <span class="original-price">$60.00</span>
      <span class="rating">4.5 (2,314)</span>
      <span class="sold-count">5.2K sold</span>
    </li>
    """,
]


def check(label: str, condition: bool, detail: str = "") -> bool:
    status = "PASS" if condition else "FAIL"
    print(f"  [{status}] {label}{(' -> ' + detail) if detail else ''}")
    return condition


def check_security_detection() -> tuple[bool | None, str]:
    """验证"安全验证检测"不会误判。

    复现的问题：TikTok 验证通过后有时会把节点留在 DOM 里，
    旧逻辑只看 innerText，于是程序以为还没通过、一直干等。
    这里要求：隐藏节点不算验证，可见遮罩才算。
    """
    try:
        from scraper import TikTokRankingScraper

        # 自检用临时会话，避免依赖 .browser_profile 的读写权限 / 占用情况；
        # 同时关闭 CDP 模式（自检不连接真人浏览器）
        scraper = TikTokRankingScraper(headless=True, use_profile=False, cdp_url="")
        scraper._start_browser()
    except Exception as exc:  # 浏览器起不来就跳过，不影响其它自检项
        return None, f"跳过（无法启动浏览器：{exc}）"

    try:
        page = scraper._context.new_page()

        page.set_content(
            "<html><head><title>Best Sellers</title></head><body>"
            '<div style="display:none">Verify to continue</div>'
            '<div style="width:900px;height:600px">products</div>'
            "</body></html>"
        )
        hidden_hit = scraper._security_check_signature(page)

        page.set_content(
            "<html><head><title>Best Sellers</title></head><body>"
            '<div style="position:fixed;inset:0;width:100%;height:100%">'
            "Verify to continue: Drag the puzzle piece into place</div>"
            "</body></html>"
        )
        visible_hit = scraper._security_check_signature(page)

        page.set_content("<html><head><title>Security Check</title></head><body><p>hi</p></body></html>")
        title_hit = scraper._security_check_signature(page)

        ok = hidden_hit is None and bool(visible_hit) and bool(title_hit)
        detail = f"隐藏节点={hidden_hit!r} 可见遮罩={visible_hit!r} 标题={title_hit!r}"
        return ok, detail
    finally:
        scraper.close()


def check_auto_escalation() -> tuple[bool, str]:
    """验证：无头模式被验证拦住时，会自动切换成可见浏览器并等待人工通过。"""
    from scraper import TikTokRankingScraper

    class DummyPage:
        pass

    class FakeScraper(TikTokRankingScraper):
        """不真启动浏览器，只走 _prepare_verified_page 的分支逻辑。"""

        def __init__(self) -> None:
            super().__init__(headless=True, use_profile=False, cdp_url="")
            self.starts = 0

        def _start_browser(self) -> None:
            self.starts += 1
            self._context = object()

        def _open_page(self):
            return DummyPage()

        def _save_debug_artifacts(self, page, tag) -> None:
            return None

        def _security_check_signature(self, page):
            # 无头时被拦；切换成可见浏览器后放行
            return "fake-security-check" if self.headless else None

    scraper = FakeScraper()
    scraper._start_browser()  # 模拟 run() 的正常启动流程
    page = scraper._prepare_verified_page()
    ok = (
        page is not None
        and scraper.headless is False
        and scraper.wait_for_verify is True
        and scraper.starts >= 2
    )
    detail = (
        f"启动次数={scraper.starts} headless={scraper.headless} "
        f"wait_for_verify={scraper.wait_for_verify}"
    )
    return ok, detail


def main() -> int:
    failures = 0

    print("1. 数值解析工具")
    failures += not check("to_count('1.2K') == 1200", to_count("1.2K") == 1200)
    failures += not check("to_count('3.4万') == 34000", to_count("3.4万") == 34000)
    failures += not check(
        "extract_prices('$23.99 ... $39.99')",
        extract_prices("$23.99 was $39.99") == [23.99, 39.99],
    )
    sold_value, sold_raw = extract_sold("4.7★ 12.5K sold")
    failures += not check("extract_sold('12.5K sold')", sold_value == 12500, f"{sold_value} / {sold_raw!r}")
    failures += not check("extract_sold('Sold 3.4K')", extract_sold("Sold 3.4K")[0] == 3400)
    failures += not check("extract_rating('4.7★')", extract_rating("4.7★") == 4.7)
    failures += not check("extract_rating('4.5 (2,314)')", extract_rating("4.5 (2,314)") == 4.5)
    failures += not check(
        "extract_discount 反算",
        extract_discount("", 42.0, 60.0) == 30.0,
        str(extract_discount("", 42.0, 60.0)),
    )

    print("\n2. 卡片解析")
    records = []
    for index, html in enumerate(CARDS):
        record = build_product_record(
            card_html=html,
            base_url=config.URL,
            keywords=config.PRODUCT_URL_KEYWORDS,
            index=index,
        )
        records.append(record)
        print(
            f"  #{index + 1} {record['product_name'][:40]!r} | 店铺={record['shop_name']!r} | "
            f"现价={record['current_price']} | 原价={record['original_price']} | "
            f"折扣={record['discount']} | 销量={record['sold_count']} | 评分={record['rating']}"
        )

    first = records[0]
    failures += not check("商品名解析", first["product_name"].startswith("Wireless Earbuds"), first["product_name"])
    failures += not check("店铺解析", first["shop_name"] == "SoundCore Official Store", first["shop_name"])
    failures += not check("现价解析", first["current_price"] == 23.99, str(first["current_price"]))
    failures += not check("原价识别（del 标签）", first["original_price"] == 39.99, str(first["original_price"]))
    failures += not check("销量解析", first["sold_count"] == 12500, str(first["sold_count"]))
    failures += not check("评分解析", first["rating"] == 4.7, str(first["rating"]))
    failures += not check("商品链接", first["product_url"].endswith("/us/view/product/1729384756"), first["product_url"])
    failures += not check("图片链接", first["image_url"].endswith("1.jpg"), first["image_url"])

    second = records[1]
    failures += not check(
        "懒加载图片（data-src）",
        second["image_url"].endswith("2.jpg"),
        second["image_url"],
    )
    failures += not check("绝对值 URL 原样保留", second["product_url"].endswith("/us/product/99887766"), second["product_url"])
    failures += not check("'Sold 3.4K' 解析", second["sold_count"] == 3400, str(second["sold_count"]))
    failures += not check("单价格场景", second["current_price"] == 15.49 and second["original_price"] is None)

    third = records[2]
    failures += not check("srcset 选最大图", third["image_url"].endswith("3.jpg"), third["image_url"])
    failures += not check("original-price class 识别", third["original_price"] == 60.0, str(third["original_price"]))
    failures += not check("折扣反算 30%", third["discount"] == 30.0, str(third["discount"]))

    print("\n3. 去重与排名")
    duplicated = records + [dict(records[0])]
    unique = assign_ranks(dedupe_records(duplicated))
    failures += not check("去重后 3 条", len(unique) == 3, str(len(unique)))
    failures += not check("排名 1..3", [r["rank"] for r in unique] == [1, 2, 3])
    failures += not check("有效性判定", all(is_meaningful(r) for r in unique))

    print("\n4. Excel 导出")
    config.EMBED_PRODUCT_IMAGES = False  # 自检用假图片链接，不实际下载
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / config.EXCEL_FILENAME
        try:
            export_to_excel(unique, out)
        except Exception as exc:  # pragma: no cover
            print(f"  [FAIL] 导出异常：{exc}")
            failures += 1
        else:
            from openpyxl import load_workbook

            workbook = load_workbook(out)
            sheet = workbook["Ranking"]
            header = [cell.value for cell in sheet[1]]
            expected = [header_text for header_text, _field, _width in config.EXCEL_COLUMNS]
            failures += not check("表头正确", header == expected, str(header))
            failures += not check("表头共 5 列", len(header) == 5, str(len(header)))
            failures += not check("行数 = 3 + 表头", sheet.max_row == 4, str(sheet.max_row))
            failures += not check(
                "第一行商品名称",
                sheet.cell(row=2, column=1).value == unique[0]["product_name"],
                str(sheet.cell(row=2, column=1).value)[:40],
            )
            failures += not check(
                "销量写入为数字",
                isinstance(sheet.cell(row=2, column=2).value, int),
            )
            failures += not check("存在『抓取说明』表", "抓取说明" in workbook.sheetnames)

    print("\n4b. 图片内嵌（离线，PIL 生成假图）")
    from openpyxl import Workbook as _Workbook
    from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor
    from excel_export import _embed_image

    import io as _io
    from PIL import Image as _PILImage

    _buf = _io.BytesIO()
    _PILImage.new("RGB", (300, 200), (200, 30, 30)).save(_buf, format="PNG")
    _wb = _Workbook()
    _ws = _wb.active
    _ws.column_dimensions["D"].width = 30
    _ws.row_dimensions[2].height = 60
    _ok = _embed_image(_ws, "D2", "https://example.com/x.png", _buf.getvalue(), 60)
    failures += not check("离线图片内嵌成功", _ok)
    failures += not check(
        "图片对象已添加", len(_ws._images) == 1, str(len(_ws._images))
    )
    failures += not check(
        "锚点为居中类型", isinstance(_ws._images[0].anchor, OneCellAnchor)
    )
    _cell_d2 = _ws["D2"]
    failures += not check(
        "内嵌后无超链接/无文本",
        _cell_d2.hyperlink is None and _cell_d2.value is None,
    )

    print("\n5. 类别映射")
    failures += not check(
        "中文名解析（美妆个护）",
        (resolve_category("美妆个护") or {}).get("slug") == "beauty-personal-care",
        str((resolve_category("美妆个护") or {}).get("url")),
    )
    failures += not check(
        "英文名解析（大小写不敏感）",
        (resolve_category("beauty & personal care") or {}).get("category_id") == "601450",
    )
    failures += not check(
        "slug 解析",
        (resolve_category("phones-electronics") or {}).get("name_zh") == "手机和电子产品",
    )
    failures += not check(
        "类目页 URL 生成",
        category_page_url("beauty-personal-care")
        == "https://shop.tiktok.com/us/c/beauty-personal-care/601450",
        category_page_url("beauty-personal-care"),
    )
    failures += not check(
        "未知类别返回 None",
        resolve_category("不存在的类别") is None,
    )
    failures += not check("类别列表非空", len(list_categories()) >= 10)

    print("\n6. PDP 类别解析")
    pdp_html_sample = (
        '<html><body><script type="application/json" id="__MODERN_ROUTER_DATA__">'
        '{"loaderData":{"(region)/pdp/(product_name_slug$)/(product_id)/page":'
        '{"page_config":{"global_data":{"product_info":{"categories":['
        '{"category_id":"601450","level":1,"is_leaf":false,"parent_id":"0",'
        '"category_name":"Beauty &amp; Personal Care"},'
        '{"category_id":"848776","level":2,"is_leaf":false,"parent_id":"601450",'
        '"category_name":"Skincare"}]}}}}}}'
        "</script></body></html>"
    )
    pdp_cat = extract_category_from_pdp(pdp_html_sample)
    failures += not check(
        "PDP 一级类别解析（含 HTML 实体）",
        bool(pdp_cat)
        and pdp_cat["category_id"] == "601450"
        and pdp_cat["category_name"] == "Beauty & Personal Care",
        str(pdp_cat),
    )
    failures += not check(
        "PDP 无分类返回 None",
        extract_category_from_pdp("<html><body>no data</body></html>") is None,
    )
    failures += not check(
        "PDP 空输入返回 None",
        extract_category_from_pdp("") is None,
    )

    print("\n7. PDP 详情介绍解析")
    # SSR JSON（components_map 里完整 description）优先
    ssr_desc_html = (
        '<html><body><script type="application/json" id="__MODERN_ROUTER_DATA__">'
        '{"loaderData":{"(region)/pdp/(product_name_slug$)/(product_id)/page":'
        '{"page_config":{"components_map":[{"component_data":{"product_info":'
        '{"product_model":{"description":"'
        '[{\\"type\\":\\"text\\",\\"text\\":\\"Radiant Glow Routine Set\\"},'
        '{\\"type\\":\\"ul\\",\\"content\\":[\\"line one\\",\\"line two\\"]},'
        '{\\"type\\":\\"image\\",\\"image\\":{}}]'
        '"}}}}]}}}}'
        "</script></body></html>"
    )
    ssr_desc = extract_product_description(ssr_desc_html)
    failures += not check(
        "SSR 描述提取（text + 列表，跳过图片）",
        ssr_desc == "Radiant Glow Routine Set\n・line one\n・line two",
        repr(ssr_desc),
    )
    # DOM 兜底
    desc_html = (
        '<div class="product-detail-section">'
        "<h3>Product description</h3>"
        "<div class=\"desc-body\">This is a great moisturizer for sensitive skin."
        " Plant-based hero ingredient. Calms redness.</div>"
        "</div>"
        '<div class="product-detail-section">'
        "<h3>Safety &amp; compliance</h3>"
        "<div>Some compliance text.</div>"
        "</div>"
    )
    desc = extract_product_description(desc_html)
    failures += not check(
        "详情介绍正文提取",
        "great moisturizer" in desc and "compliance" not in desc,
        desc[:80],
    )
    failures += not check(
        "无描述返回空字符串",
        extract_product_description("<html><body><p>no desc</p></body></html>") == "",
    )
    failures += not check(
        "空输入返回空字符串",
        extract_product_description("") == "",
    )

    print("\n8. 安全验证检测（不联网，仅本地浏览器）")
    ok, detail = check_security_detection()
    if ok is None:
        print(f"  [SKIP] {detail}")
    else:
        failures += not check("隐藏验证层不误判 / 可见遮罩与标题能识别", ok, detail)

    print("\n9. 验证拦截时的自动升级（无头 → 可见窗口）")
    ok, detail = check_auto_escalation()
    failures += not check("被拦后自动切到可见浏览器并进入等待", ok, detail)

    print()
    if failures:
        print(f"自检结束：{failures} 项失败，请检查上面的输出。")
        return 1
    print("自检结束：全部通过 ✅  可以执行 python main.py 了。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
