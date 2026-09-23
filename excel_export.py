"""Excel 导出模块：把商品列表写入 ``TikTok_US_Ranking.xlsx``。"""

from __future__ import annotations

import io
import logging
import math
import re
import time
import urllib.request
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.worksheet import Worksheet

import config

log = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="111827")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LINK_FONT = Font(color="0563C1", underline="single")

#: 下载图片时用的 User-Agent（部分 CDN 会拦截无 UA 请求）
_DOWNLOAD_UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
)

#: 内嵌商品图的边长上限（像素），越大越清晰，但 Excel 文件也越大。
#: 220 以上时源图（通常 2000px）缩到该尺寸后放大查看仍锐利。
EMBED_IMAGE_SIZE = 260

#: 图片太小（小于该边长）时，尝试把 URL 的 resize 参数改大重新下载
RESIZE_UPGRADE_THRESHOLD = 400
RESIZE_UPGRADE_TARGET = 2000


def _safe_sheet_name(name: str) -> str:
    name = re.sub(r"[\[\]:*?/\\]", "-", name or "Ranking")
    return name[:31] or "Ranking"


#: Excel 单元格最大字符数（超出会写入失败），用于截断超长描述
MAX_CELL_CHARS = 32000


def _cell_value(field: str, value: Any) -> Any:
    """把字段值转成适合写进单元格的类型。"""
    if value is None or value == "":
        return None
    if field in ("current_price", "original_price"):
        return round(float(value), 2)
    if field in ("discount", "rating"):
        return round(float(value), 1)
    if field in ("sold_count", "rank"):
        return int(value)
    text = str(value)
    if len(text) > MAX_CELL_CHARS:
        text = text[:MAX_CELL_CHARS]
    return text


def _download_image_bytes(url: str, timeout: int = 15) -> bytes | None:
    """下载图片字节；小图且 URL 带 resize 参数时，尝试换大尺寸重新下载。

    返回 bytes；下载失败返回 None。
    """
    # TikTok 图片 CDN 会做防盗链检查：必须带 Referer（来自 TikTok Shop 页面），
    # 否则请求会被 403 拒绝 —— 这是"链接能打开、程序却抓不到图"的主因。
    headers = {
        "User-Agent": _DOWNLOAD_UA,
        "Referer": "https://shop.tiktok.com/",
        "Accept": "image/avif,image/webp,image/apng,image/svg+xml,image/*,*/*;q=0.8",
    }

    def fetch(target: str, attempt: int = 1) -> bytes | None:
        try:
            request = urllib.request.Request(target, headers=headers)
            raw = urllib.request.urlopen(request, timeout=timeout).read()
            return raw if raw and len(raw) >= 100 else None
        except Exception as exc:
            if attempt < 3:
                time.sleep(0.8 * attempt)
                return fetch(target, attempt + 1)
            log.warning(
                "图片下载失败（已重试 3 次）：%s -> %s",
                target[:110],
                exc,
            )
            return None

    raw = fetch(url)
    if not raw:
        return None
    try:
        from PIL import Image

        with Image.open(io.BytesIO(raw)) as img:
            width, height = img.size
    except Exception:
        return raw
    # 图片太小且 URL 可改尺寸 → 用更大 resize 再试一次
    if max(width, height) < RESIZE_UPGRADE_THRESHOLD and re.search(
        r"resize-png:\d+:\d+", url
    ):
        bigger_url = re.sub(
            r"resize-png:\d+:\d+",
            f"resize-png:{RESIZE_UPGRADE_TARGET}:{RESIZE_UPGRADE_TARGET}",
            url,
        )
        if bigger_url != url:
            raw_bigger = fetch(bigger_url)
            if raw_bigger:
                return raw_bigger
    return raw


def _embed_image(
    worksheet: Worksheet,
    anchor: str,
    url: str,
    raw: bytes | None,
    row_height_pt: float,
    max_size: int = EMBED_IMAGE_SIZE,
) -> bool:
    """把已下载的商品图片内嵌到指定单元格（等比缩放到 max_size、水平垂直居中）。

    内嵌成功返回 True，且不设置超链接（单元格只放图片）；
    任何失败返回 False（调用方退回链接文本 + 超链接）。
    """
    try:
        from PIL import Image

        if not raw:
            return False
        with Image.open(io.BytesIO(raw)) as img:
            img.thumbnail((max_size, max_size), Image.Resampling.LANCZOS)
            width, height = img.size
            buffer = io.BytesIO()
            # 统一转成 PNG 内嵌（openpyxl 对 PNG 支持最稳，webp 也能转）
            img.convert("RGB").save(buffer, format="PNG")
        buffer.seek(0)

        match = re.match(r"([A-Z]+)(\d+)", anchor)
        if not match:
            return False
        col_letter, row_idx = match.group(1), int(match.group(2))
        col_idx = column_index_from_string(col_letter)

        # 单元格像素尺寸（Excel 近似换算：列宽字符 ≈ 7px/字符 + 5px）
        col_width = worksheet.column_dimensions[col_letter].width or 10
        cell_w = col_width * 7 + 5
        cell_h = row_height_pt * 4 / 3  # 磅 → 像素

        off_x = max(0, int((cell_w - width) // 2))
        off_y = max(0, int((cell_h - height) // 2))

        marker = AnchorMarker(
            col=col_idx - 1,
            colOff=pixels_to_EMU(off_x),
            row=row_idx - 1,
            rowOff=pixels_to_EMU(off_y),
        )
        xl_image = XLImage(buffer)
        xl_image.width = width
        xl_image.height = height
        xl_image.anchor = OneCellAnchor(
            _from=marker,
            ext=XDRPositiveSize2D(pixels_to_EMU(width), pixels_to_EMU(height)),
        )
        worksheet.add_image(xl_image)
        return True
    except Exception as exc:  # noqa: BLE001 - 任何失败都退化为链接
        log.warning("商品图片内嵌失败（改用链接）：%s -> %s", url[:90], exc)
        return False


def _style_sheet(
    worksheet: Worksheet,
    columns: list[tuple[str, str, int]],
    row_count: int,
) -> None:
    for col_index, (_header, field, width) in enumerate(columns, start=1):
        letter = get_column_letter(col_index)
        worksheet.column_dimensions[letter].width = width

        header_cell = worksheet.cell(row=1, column=col_index)
        header_cell.fill = HEADER_FILL
        header_cell.font = HEADER_FONT
        header_cell.alignment = Alignment(horizontal="center", vertical="center")
        header_cell.border = BORDER

        for row in range(2, row_count + 2):
            cell = worksheet.cell(row=row, column=col_index)
            cell.border = BORDER

            if field in ("rank", "sold_count", "rating"):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif field in ("current_price", "original_price", "discount"):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif field == "product_name":
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            elif field == "description":
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="center")

            if field in ("current_price", "original_price"):
                cell.number_format = '"$"#,##0.00'
            elif field == "discount":
                cell.number_format = '0.0"%"'
            elif field == "rating":
                cell.number_format = "0.0"
            elif field == "sold_count":
                cell.number_format = "#,##0"
            elif field in ("product_url", "image_url") and cell.value:
                cell.hyperlink = str(cell.value)
                cell.font = LINK_FONT

    worksheet.row_dimensions[1].height = 22
    worksheet.freeze_panes = "A2"
    if row_count > 0:
        worksheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{row_count + 1}"


def export_to_excel(
    products: Iterable[dict[str, Any]],
    output_path: str | Path,
    sheet_name: str = "Ranking",
    category_label: str | None = None,
) -> Path:
    """把商品记录写入 Excel，返回最终文件路径。"""
    products = list(products)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    columns = config.EXCEL_COLUMNS
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = _safe_sheet_name(sheet_name)

    # 表头 + 列宽（先设列宽，图片居中计算需要）
    for col_index, (header, _field, width) in enumerate(columns, start=1):
        worksheet.cell(row=1, column=col_index, value=header)
        worksheet.column_dimensions[get_column_letter(col_index)].width = width

    # 第一遍：下载所有商品图、记录内嵌尺寸（只下一次网）
    image_cache: dict[str, bytes | None] = {}
    row_image_sizes: dict[int, tuple[int, int]] = {}
    if config.EMBED_PRODUCT_IMAGES:
        from PIL import Image as PILImage

        for row_offset, product in enumerate(products, start=2):
            url = str(product.get("image_url") or "")
            if not url:
                continue
            raw = image_cache.setdefault(url, _download_image_bytes(url))
            if raw:
                try:
                    with PILImage.open(io.BytesIO(raw)) as img:
                        img.thumbnail(
                            (EMBED_IMAGE_SIZE, EMBED_IMAGE_SIZE),
                            PILImage.Resampling.LANCZOS,
                        )
                        row_image_sizes[row_offset] = img.size
                except Exception:
                    row_image_sizes.pop(row_offset, None)
                    image_cache[url] = None

    # 第二遍：写单元格、按图片与介绍自适应行高、内嵌图片
    for row_offset, product in enumerate(products, start=2):
        desc = str(product.get("description") or "")
        desc_lines = max(1, math.ceil(len(desc) / 60))
        desc_pt = 24 + desc_lines * 14
        img_size = row_image_sizes.get(row_offset)
        img_pt = (img_size[1] * 0.75 + 8) if img_size else 0
        row_height = max(60, min(320, desc_pt), img_pt)
        worksheet.row_dimensions[row_offset].height = row_height

        for col_index, (_header, field, _width) in enumerate(columns, start=1):
            letter = get_column_letter(col_index)
            value = product.get(field)
            if (
                field == "image_url"
                and value
                and config.EMBED_PRODUCT_IMAGES
                and img_size
                and _embed_image(
                    worksheet,
                    f"{letter}{row_offset}",
                    str(value),
                    image_cache.get(str(value)),
                    row_height,
                )
            ):
                # 图片已内嵌：单元格只放图片，不加超链接
                continue
            worksheet.cell(
                row=row_offset,
                column=col_index,
                value=_cell_value(field, value),
            )

    _style_sheet(worksheet, columns, len(products))

    # 附加一张说明表，方便以后回溯本次抓取口径
    meta = workbook.create_sheet("抓取说明")
    meta_rows = [
        ("数据来源", config.URL),
        ("导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("商品数量", len(products)),
        ("排序说明", "行顺序 = 榜单页面展示顺序（第一行 = 榜单第一）"),
        ("字段说明", "商品销量解析不到时留空；商品详情介绍来自商品详情页（PDP）"),
    ]
    if category_label:
        meta_rows.insert(1, ("抓取类别", category_label))
    for row_index, (key, value) in enumerate(meta_rows, start=1):
        meta.cell(row=row_index, column=1, value=key).font = Font(bold=True)
        meta.cell(row=row_index, column=2, value=value)
    meta.column_dimensions["A"].width = 14
    meta.column_dimensions["B"].width = 80

    workbook.save(output_path)
    log.info("Excel 已保存: %s（%d 条记录）", output_path, len(products))
    return output_path


def timestamped_copy_path(
    output_dir: str | Path,
    filename: str = config.EXCEL_FILENAME,
) -> Path:
    """生成带时间戳的历史副本路径，例如 ``TikTok_US_Ranking_20260915_213000.xlsx``。"""
    stem = Path(filename).stem
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return Path(output_dir) / f"{stem}_{stamp}.xlsx"
